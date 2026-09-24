"""Build the live bracket from the published qualifying field and completed MFL scores.

No future workbook rating or unpublished game result is used. Current-season
qualification stays owned by the existing forecast, with postseason progression here.
"""
from pathlib import Path
import argparse, csv, html, io, json, math, re, urllib.request
import numpy as np
import openpyxl
import probability as prob

ROOT=Path(__file__).resolve().parents[1]
BOOKS={'ADL':'1iu7oJUQ8IEhHDTp5RiArK7oTD4-DmjTbI1xtOwuwRBI','FAFL':'1yWEzFx8hKhhlTQ47gacHSQXmZtPsX9k7hB2s6D6-g3k'}
def fetch(url):
    with urllib.request.urlopen(url,timeout=120) as r:return r.read()
def ratings(book,league,year,week):
    rows=list(book['Graphs' if league=='ADL' else str(year)].values)
    label=f'{year%100}W{week}Elo' if league=='ADL' else f'W{week}Elo'
    col=next(i for i,v in enumerate(rows[0]) if v==label)
    team_rows=rows[3:35] if league=='ADL' else rows[1:33]
    values={}
    missing=[]
    for row in team_rows:
        team=str(row[0]).strip() if row and row[0] is not None else '<unknown team>'
        if not row or row[0] is None or col>=len(row) or row[col] is None:
            missing.append(team)
            continue
        try:
            value=float(row[col])
        except (TypeError,ValueError):
            missing.append(team)
            continue
        if math.isfinite(value):
            values[team]=value
        else:
            missing.append(team)
    if len(values)!=32:
        raise ValueError(f'Incomplete Elo ratings for {label}: found {len(values)} of 32 teams; missing {", ".join(missing)}')
    return values

def internal_ratings(path,year,first_week,last_week):
    rows=list(csv.DictReader(path.open(encoding='utf8')));result={}
    for week in range(first_week,last_week+1):
        values={r['franchise_name']:float(r['elo']) for r in rows if int(r['season'])==year and int(r['week'])==week}
        if len(values)!=32 or any(not math.isfinite(v) for v in values.values()):raise ValueError(f'Internal Elo cache is incomplete for Week {week}')
        result[week]=values
    return result

def make_games(field,completed,scoreweeks,elo,scale):
    prob.scale=scale
    teams={html.unescape(t['name']):dict(t,name=html.unescape(t['name'])) for ts in field.values() for t in ts}
    if len(teams)!=32:raise ValueError('Expected 32 distinct teams')
    reg=min(completed,12);ap={n:round(float(t['apPct'])*31*reg*2)/2 for n,t in teams.items()};pf={n:float(t.get('pointsTotal',float(t['ppg'])*reg)) for n,t in teams.items()};pot={n:float(t['pot'])*reg for n,t in teams.items()}
    result=[];ranks={n:t['seed'] for n,t in teams.items()}
    states={}
    for conf,ts in field.items():
        names=[html.unescape(t['name']) for t in sorted(ts,key=lambda t:t['seed'])]
        states[conf]={'playoff':names[:7],'survivors':names[:7],'losers':[], 'ladder':[names[7:9],names[9:11],names[11:13],names[13:]],'wc':[names[1:2]+names[6:7],names[2:3]+names[5:6],names[3:5]],'bye':names[0],'series':None}
    def ranking(ns):return sorted(ns,key=lambda n:(-ap[n],-pf[n],-pot[n],n))
    def points(n,w):return scoreweeks[w][n]['points']
    def winner(g,total=False):return min(g['teams'],key=lambda t:(-(t['aggregate'] if total else t['points']),t['seed']))['name']
    def add(w,conf,cat,title,ns,first=None,series=None):
        es=elo[min(completed,w-1)];ts=[]
        for n in ns:
            t=teams[n];bank=points(n,first) if first and first<w else 0
            now=points(n,w) if w<=completed else None
            ts.append(dict(name=n,logo=t['logo'],seed=ranks[n],record=t['record'],allPlayPct=ap[n]/(31*min(completed,w-1))*100,elo=es[n],points=now,aggregate=now+bank if now is not None else None,bankedPoints=bank))
        legs=2 if first==w else 1
        odds=prob.two_leg(ts) if legs==2 and len(ts)>1 else prob.softmax([prob.c*t['elo']+t['bankedPoints']/scale for t in ts])
        g=dict(week=w,conference=conf,category=cat,title=title,teams=ts,remainingWeeks=legs,probabilityLabel='win prob')
        if series and set(series['priorWins'])==set(ns):
            a,b=ns;d=series['priorWins'][a]-series['priorWins'][b];lead=series['priorPoints'][a]-series['priorPoints'][b]
            threshold=min(0,-lead) if d==1 else max(0,-lead) if d==-1 else 0
            p=1 if d>=2 else 0 if d<=-2 else prob.sigmoid(prob.c*(es[a]-es[b])-threshold/scale)
            odds=[p,1-p];g['series']=series.copy()
        for t,p in zip(ts,odds):t['odds']=float(p)*100
        if w<=completed and series:
            wins=series['priorWins'].copy();pts=series['priorPoints'].copy();wins[winner(g)]+=1
            for n in ns:pts[n]+=points(n,w)
            g['seriesWinner']=min(ns,key=lambda n:(-wins[n],-pts[n],ranks[n]));g['series']['winsAfter']=wins
        result.append(g);return g
    for w in range(13,min(17,completed+1)+1 if completed>=12 else 14):
        for conf,st in states.items():
            if w>=15:
                for offset,group in [(0,st['survivors']),(4,st['losers']),(7,sum(st['ladder'],[]))]:
                    for i,n in enumerate(ranking(group),offset+1):ranks[n]=i
            elif w==14:
                for offset,group in [(0,st['playoff']),(7,sum(st['ladder'],[]))]:
                    for i,n in enumerate(ranking(group),offset+1):ranks[n]=i
            if w<=14:
                add(w,conf,'championship','Bye',[st['bye']])
                matches=[add(w,conf,'championship','Wild Card Round',ns,13) for ns in st['wc']]
                if w==14 and w<=completed:
                    st['survivors']=[st['bye']]+[winner(g,True) for g in matches];st['losers']=[n for n in st['playoff'] if n not in st['survivors']]
            if w==15:
                ns=ranking(st['survivors']);matches=[add(w,conf,'championship','Divisional Round',pair) for pair in [[ns[0],ns[3]],[ns[1],ns[2]]]]
                if w<=completed:st['finalists']=[winner(g) for g in matches];st['third']=[n for n in ns if n not in st['finalists']]
            if w==16:
                title=add(w,conf,'championship',conf+' Championship',st['finalists']);third=add(w,conf,'placement','3rd Place Game',st['third'])
                if w<=completed:st['finish']=[winner(title)]+[n for n in st['finalists'] if n!=winner(title)]+[winner(third)]+[n for n in st['third'] if n!=winner(third)]
            if w in [15,16]:
                fifth=add(w,conf,'placement','5th Place Game',st['losers'],15)
                if w==16 and w<=completed:st['finish'] += [t['name'] for t in sorted(fifth['teams'],key=lambda t:(-t['aggregate'],t['seed']))]
            gs=[]
            for i,ns in enumerate(st['ladder']):
                series=st['series'] if i==0 and st['series'] and set(st['series']['priorWins'])==set(ns) else None
                if i==0 and not series:series={'startWeek':w,'priorWins':dict.fromkeys(ns,0),'priorPoints':dict.fromkeys(ns,0)}
                gs.append(add(w,conf,'ladder',f'Consolation Game {i+1}',ns,series=series))
            if w<=completed:
                a=gs[0].get('seriesWinner',winner(gs[0]));b=next(n for n in st['ladder'][0] if n!=a);c=winner(gs[1]);d=next(n for n in st['ladder'][1] if n!=c);e=winner(gs[2]);f=next(n for n in st['ladder'][2] if n!=e);g=winner(gs[3]);bottom=[n for n in st['ladder'][3] if n!=g]
                promote=points(c,w)>points(b,w)
                st['ladder']=[[a,c if promote else b],[b if promote else c,e],[d,g],[f]+bottom]
                old=gs[0];st['series']={'startWeek':old['series']['startWeek'],'priorWins':old['series']['winsAfter'],'priorPoints':{t['name']:old['series']['priorPoints'][t['name']]+t['points'] for t in old['teams']}}
        if w==17:
            for i in range(7):add(w,'ADL','championship' if i==0 else 'placement','Super Bowl' if i==0 else f'{i+1}{"nd" if i==1 else "rd" if i==2 else "th"} Place', [states[c]['finish'][i] for c in ['NFC','AFC']])
        if w<=completed:
            vals=scoreweeks[w]
            for n in teams:
                p=vals[n]['points'];ap[n]+=sum(p>v['points'] for m,v in vals.items() if m!=n)+.5*sum(p==v['points'] for m,v in vals.items() if m!=n);pf[n]+=p;pot[n]+=vals[n]['potential']
    return result

def main():
    p=argparse.ArgumentParser();p.add_argument('--league',choices=BOOKS,required=True);p.add_argument('--book');a=p.parse_args();league=a.league
    year=2026;out=ROOT/'docs/playoff-picture';page=(out/'index.html').read_text(encoding='utf8');field=json.JSONDecoder().raw_decode(page.split('const data=',1)[1].lstrip())[0]
    if league=='ADL':completed=int(next(csv.DictReader((ROOT/'data/playoff_picture_metadata.csv').open()))['through_week'])
    else:completed=int(json.loads((ROOT/'data/current_forecast.json').read_text())['through_week'])
    if league=='FAFL' and not a.book:
        elo=internal_ratings(ROOT/'data/elo_ratings.csv',year,min(completed,12),completed)
    else:
        raw=Path(a.book).read_bytes() if a.book else fetch(f'https://docs.google.com/spreadsheets/d/{BOOKS[league]}/export?format=xlsx')
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
        elo={w:ratings(book,league,year,w) for w in range(min(completed,12),completed+1)}
        if league=='ADL':
            available=sum(isinstance(v,(int,float)) for v in next(book['Data'].iter_rows(min_row=37,max_row=37,min_col=167,max_col=183,values_only=True)))
        else:
            available=sum(isinstance(v,(int,float)) for v in next(book[str(year)].iter_rows(min_row=37,max_row=37,min_col=2,max_col=18,values_only=True)))
        if available<completed:raise ValueError(f'Elo workbook is only current through Week {available}; refusing future calculated ratings for Week {completed}')
    scoreweeks={};lid='60206' if league=='ADL' else '22686'
    if completed>12:
        if league=='FAFL':
            rows=list(csv.DictReader((ROOT/'data/current_weekly.csv').open(encoding='utf8')))
            names={str(t['franchise_id']).zfill(4):html.unescape(t['name']) for ts in field.values() for t in ts}
            for w in range(13,completed+1):
                scoreweeks[w]={names[str(r['franchise_id']).zfill(4)]:{'points':float(r['points']),'potential':float(r['potential'])} for r in rows if int(r['week'])==w}
                if len(scoreweeks[w])!=32:raise ValueError(f'Incomplete shared Week {w} scores')
        else:
            meta=json.loads(fetch(f'https://api.myfantasyleague.com/{year}/export?TYPE=league&L={lid}&JSON=1'))['league']['franchises']['franchise'];names={f['id']:f['name'] for f in meta}
            for w in range(13,completed+1):
                data=json.loads(fetch(f'https://api.myfantasyleague.com/{year}/export?TYPE=weeklyResults&L={lid}&W={w}&JSON=1'))['weeklyResults'];fs=[f for m in data['matchup'] for f in m['franchise']];fs+=data.get('franchise',[])
                scoreweeks[w]={names[f['id']]:{'points':float(f['score']),'potential':float(f['opt_pts'])} for f in fs}
                if len(scoreweeks[w])!=32:raise ValueError(f'Incomplete Week {w} scores')
    scale=30.625025877938 if league=='ADL' else 31.57649140329526
    games=make_games(field,completed,scoreweeks,elo,scale)
    for g in games:
        if g['conference']=='ADL':g['conference']=league
    data=dict(league=league,season=year,completedWeek=completed,games=games,probabilityModel=dict(logisticPointsScale=scale),eloThroughWeek=completed)
    template=(ROOT/'scripts/templates/bracket.html').read_text(encoding='utf8');text=template.replace('__DATA__',json.dumps(data,allow_nan=False).replace('</','<\\/')).replace('__LEAGUE__',league).replace('__SEASON__',str(year))
    archive=f'bracket-week-{completed+1:02}.html'
    (out/archive).write_text(text,encoding='utf8')
    for report in out.glob('*.html'):
        if report.name.startswith('bracket'):continue
        content=report.read_text(encoding='utf8')
        if 'src="bracket.html"' in content:report.write_text(content.replace('src="bracket.html"',f'src="{archive}"'),encoding='utf8')
    (out/'bracket.html').write_text(text,encoding='utf8');(out/'bracket-data.json').write_text(json.dumps(data,indent=2,allow_nan=False),encoding='utf8')
    print(f'{league}: built {len(games)} games from completed Week {completed}; points scale {scale:.3f}')
if __name__=='__main__':main()
